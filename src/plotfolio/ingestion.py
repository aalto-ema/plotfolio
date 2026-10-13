"""Previewable CSV/TSV and Excel ingestion with explicit column metadata."""

import csv
from pathlib import Path

from openpyxl import load_workbook

ROLES = ["Unassigned", "X", "Y", "Y error", "X error", "Group", "Ignore"]


def table_dataset(path, cells, header_row=1, has_header=True, sheet=""):
    if header_row < 1:
        raise ValueError("First header/data row must start at 1.")
    cells = list(cells)[header_row - 1 :]
    cells = [
        list(row) for row in cells if any(value is not None and str(value).strip() for value in row)
    ]
    if not cells:
        raise ValueError("No rows after the selected start row.")
    count = max(len(row) for row in cells)
    for row in cells:
        row.extend([None] * (count - len(row)))
    if has_header:
        headers = [str(value).strip() if value is not None else "" for value in cells.pop(0)]
    else:
        headers = [f"Column {index + 1}" for index in range(count)]
    if not headers or any(not header for header in headers) or len(set(headers)) != len(headers):
        raise ValueError("Choose a row with unique, nonempty column headers, or turn headers off.")
    if not cells:
        raise ValueError("The table contains no data rows.")
    rows = [[str(value) if value is not None else "" for value in row] for row in cells]
    return dict(
        name=Path(path).name + (f" · {sheet}" if sheet else ""),
        headers=headers,
        rows=rows,
        source=str(Path(path).resolve()),
        column_metadata={
            header: {"role": "Unassigned", "unit": "", "label": header} for header in headers
        },
        import_options={"header_row": header_row, "has_header": has_header, "sheet": sheet},
    )


def sheet_names(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return book.sheetnames
    finally:
        book.close()


def read_tables(path, delimiter="Auto", header_row=1, has_header=True, sheet="All sheets"):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            names = book.sheetnames if sheet == "All sheets" else [sheet]
            return [
                table_dataset(
                    path, book[name].iter_rows(values_only=True), header_row, has_header, name
                )
                for name in names
            ]
        finally:
            book.close()
    separators = {"Comma": ",", "Semicolon": ";", "Tab": "\t", "Space": " "}
    with path.open(encoding="utf-8-sig", newline="") as stream:
        sample = stream.read(16384)
        stream.seek(0)
        if delimiter == "Auto":
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
            cells = list(csv.reader(stream, dialect))
        else:
            cells = list(
                csv.reader(
                    stream, delimiter=separators[delimiter], skipinitialspace=delimiter == "Space"
                )
            )
    # Unlike Excel's padded cells, CSV rows with inconsistent widths indicate an import error.
    relevant = [row for row in cells[header_row - 1 :] if row]
    if relevant and any(len(row) != len(relevant[0]) for row in relevant):
        raise ValueError("Rows have different column counts. Check delimiter and start row.")
    dataset = table_dataset(path, cells, header_row, has_header)
    dataset["import_options"]["delimiter"] = delimiter
    return [dataset]


def axis_label(dataset, column):
    info = dataset.get("column_metadata", {}).get(column, {})
    label = info.get("label") or column
    return f"{label} / {info['unit']}" if info.get("unit") else label
